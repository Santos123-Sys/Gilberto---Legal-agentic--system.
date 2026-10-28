import unittest
from io import BytesIO
from unittest.mock import MagicMock, patch
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from pypdf import PdfWriter

from document_extractor import DocumentExtractionError, extract_document_text


class DocumentExtractorTests(unittest.TestCase):
    def test_extracts_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("CONTRATO DE PRESTAÇÃO DE SERVIÇOS")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Contratante"
        table.cell(0, 1).text = "Contratada"
        document.add_paragraph("A prestação terá vigência por doze meses.")
        buffer = BytesIO()
        document.save(buffer)

        text = extract_document_text("contract.docx", buffer.getvalue())

        self.assertIn("CONTRATO DE PRESTAÇÃO DE SERVIÇOS", text)
        self.assertIn("Contratante | Contratada", text)
        self.assertIn("doze meses", text)

    def test_rejects_corrupt_docx_with_actionable_message(self):
        with self.assertRaisesRegex(DocumentExtractionError, "valid .docx"):
            extract_document_text("broken.docx", b"not a Word document")

    def test_identifies_image_only_pdf_as_needing_ocr(self):
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        buffer = BytesIO()
        writer.write(buffer)

        with self.assertRaisesRegex(DocumentExtractionError, "OCR"):
            extract_document_text("scanned.pdf", buffer.getvalue())

    def test_rejects_compressed_docx_expansion(self):
        buffer = BytesIO()
        with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"x" * (21 * 1024 * 1024))
        with self.assertRaisesRegex(DocumentExtractionError, "decompression limits"):
            extract_document_text("archive.docx", buffer.getvalue())

    def test_rejects_pdf_with_excessive_pages_before_extraction(self):
        reader = MagicMock()
        reader.is_encrypted = False
        reader.pages.__len__.return_value = 251
        with patch("document_extractor.PdfReader", return_value=reader):
            with self.assertRaisesRegex(DocumentExtractionError, "too many pages"):
                extract_document_text("contract.pdf", b"%PDF-placeholder")
        reader.pages.__iter__.assert_not_called()

    def test_rejects_excessive_plain_text(self):
        with self.assertRaisesRegex(DocumentExtractionError, "too large"):
            extract_document_text("huge.txt", b"a" * 2_000_050)

    def test_decodes_legacy_portuguese_text(self):
        text = extract_document_text("contract.txt", ("Cláusula de prestação " * 5).encode("cp1252"))
        self.assertIn("Cláusula", text)

    def test_rejects_unsupported_extension(self):
        with self.assertRaisesRegex(DocumentExtractionError, "Unsupported file type"):
            extract_document_text("contract.doc", b"content")


if __name__ == "__main__":
    unittest.main()
